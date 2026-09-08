from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from docx import Document
import io

from .analyzer import extract_skills, match_job_description

app = FastAPI(
    title="AI Resume-Job Matcher",
    description="NLP-powered resume and job description matching system"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"]
)


@app.get("/")
def home():
    return {
        "message": "AI Resume Analyzer API is running!",
        "status": "success"
    }


@app.post("/upload-resume")
async def upload_resume(file: UploadFile = File(...)):

    contents = await file.read()

    filename = file.filename.lower()

    if filename.endswith(".pdf"):

        pdf = PdfReader(io.BytesIO(contents))

        text = ""

        for page in pdf.pages:
            text += page.extract_text() or ""

        pages = len(pdf.pages)

    elif filename.endswith(".docx"):

        document = Document(io.BytesIO(contents))

        text = ""

        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

        pages = None

    else:

        return {
            "error": "Unsupported file type. Please upload a PDF or DOCX file."
        }

    skills = extract_skills(text)

    return {
        "filename": file.filename,
        "pages": pages,
        "skills": skills,
        "resume_text": text
    }
from pydantic import BaseModel


class JobDescription(BaseModel):
    resume_text: str
    job_description: str


@app.post("/match-job")
def match_job(data: JobDescription):

    result = match_job_description(
        data.resume_text,
        data.job_description
    )

    return result
@app.post("/analyze-resume")
async def analyze_resume(
    file: UploadFile = File(...),
    job_description: str = Form(...)
):

    contents = await file.read()

    filename = file.filename.lower()

    if filename.endswith(".pdf"):

        pdf = PdfReader(io.BytesIO(contents))

        text = ""

        for page in pdf.pages:
            text += page.extract_text() or ""

    elif filename.endswith(".docx"):

        document = Document(io.BytesIO(contents))

        text = ""

        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

    else:

        return {
            "error": "Unsupported file type. Please upload a PDF or DOCX file."
        }

    result = match_job_description(
        text,
        job_description
    )

    return {
        "filename": file.filename,
        "resume_text_length": len(text),
        **result
    }